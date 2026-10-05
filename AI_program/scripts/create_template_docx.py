import docx
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

def set_page_number_field(run):
    """Menyisipkan field dinamis PAGE ke dalam run Word."""
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def create_almuslim_template(output_path):
    doc = Document()
    
    # 1. Konfigurasi Gaya Dokumen (Styles)
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Times New Roman'
    font_normal.size = Pt(12)
    font_normal.color.rgb = RGBColor(0, 0, 0)
    p_format = style_normal.paragraph_format
    p_format.line_spacing = 2.0
    p_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_format.space_before = Pt(0)
    p_format.space_after = Pt(0)
    
    style_h1 = doc.styles['Heading 1']
    font_h1 = style_h1.font
    font_h1.name = 'Times New Roman'
    font_h1.size = Pt(14)
    font_h1.bold = True
    font_h1.color.rgb = RGBColor(0, 0, 0)
    h1_format = style_h1.paragraph_format
    h1_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1_format.line_spacing = 1.5
    h1_format.space_before = Pt(12)
    h1_format.space_after = Pt(12)
    
    style_h2 = doc.styles['Heading 2']
    font_h2 = style_h2.font
    font_h2.name = 'Times New Roman'
    font_h2.size = Pt(12)
    font_h2.bold = True
    font_h2.color.rgb = RGBColor(0, 0, 0)
    h2_format = style_h2.paragraph_format
    h2_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2_format.line_spacing = 1.5
    h2_format.space_before = Pt(6)
    h2_format.space_after = Pt(6)

    style_h3 = doc.styles['Heading 3']
    font_h3 = style_h3.font
    font_h3.name = 'Times New Roman'
    font_h3.size = Pt(12)
    font_h3.bold = True
    font_h3.color.rgb = RGBColor(0, 0, 0)
    h3_format = style_h3.paragraph_format
    h3_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3_format.line_spacing = 1.5
    h3_format.space_before = Pt(6)
    h3_format.space_after = Pt(6)

    def apply_section_geometry(sec):
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(4.0)
        sec.bottom_margin = Cm(3.0)
        sec.left_margin = Cm(4.0)
        sec.right_margin = Cm(3.0)

    # SECTION 1: COVER
    sec_cover = doc.sections[0]
    apply_section_geometry(sec_cover)
    sec_cover.header.is_linked_to_previous = False
    sec_cover.footer.is_linked_to_previous = False
    
    p_cover_title = doc.add_paragraph()
    p_cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_cover_title.add_run("JUDUL SKRIPSI HURUF KAPITAL CETAK TEBAL DENGAN FORMULASI SIMETRIS PIRAMIDA TERBALIK MAKSIMAL TIGA BARIS\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    p_cover_sub = doc.add_paragraph()
    p_cover_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cover_sub.paragraph_format.line_spacing = 1.0
    r = p_cover_sub.add_run("\nSKRIPSI\n\nDiajukan Untuk Memenuhi Salah Satu Syarat\nMemperoleh Gelar Sarjana Komputer (S.Kom)\nPada Program Studi Informatika\n\n\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    logo_path = os.path.abspath("ASSETS/logo_umuslim.jpg")
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(logo_path):
        p_logo.add_run().add_picture(logo_path, width=Cm(4.5))
    else:
        p_logo.add_run("[LOGO UNIVERSITAS ALMUSLIM 4.5 CM]").bold = True

    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.line_spacing = 1.0
    r = p_author.add_run("\n\nOLEH:\n\nNAMA MAHASISWA\nNPM. 15050200XX\n\n\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.line_spacing = 1.15
    r = p_inst.add_run("PROGRAM STUDI INFORMATIKA\nFAKULTAS ILMU KOMPUTER\nUNIVERSITAS ALMUSLIM\nBIREUEN\n2026")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.bold = True

    # SECTION 2: BAGIAN AWAL
    sec_prelim = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_section_geometry(sec_prelim)
    sec_prelim.header.is_linked_to_previous = False
    sec_prelim.footer.is_linked_to_previous = False
    
    pg_roman = parse_xml(r'<w:pgNumType %s w:fmt="lowerRoman" w:start="1"/>' % nsdecls('w'))
    sec_prelim._sectPr.append(pg_roman)
    
    footer_prelim = sec_prelim.footer
    p_foot_prelim = footer_prelim.paragraphs[0]
    p_foot_prelim.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = p_foot_prelim.add_run()
    r_foot.font.name = 'Times New Roman'
    r_foot.font.size = Pt(12)
    set_page_number_field(r_foot)

    doc.add_paragraph("HALAMAN PENGESAHAN", style='Heading 1')
    p_p_body = doc.add_paragraph("[Tempat tanda tangan Pembimbing I, Pembimbing II, dan Dekan FIKOM]")
    p_p_body.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()
    doc.add_paragraph("HALAMAN PERNYATAAN KEASLIAN", style='Heading 1')
    doc.add_paragraph("Saya yang bertanda tangan di bawah ini menyatakan bahwa skripsi ini adalah murni hasil karya saya sendiri...")
    
    doc.add_page_break()
    doc.add_paragraph("KATA PENGANTAR", style='Heading 1')
    doc.add_paragraph("Puji dan syukur penulis panjatkan ke hadirat Allah SWT atas rahmat dan karunia-Nya...")
    
    doc.add_page_break()
    doc.add_paragraph("ABSTRAK", style='Heading 1')
    doc.add_paragraph("Abstrak disusun dalam satu paragraf tunggal dengan jumlah 150 - 250 kata, spasi tunggal (1.0). Berisi latar belakang, tujuan, metode yang digunakan, dan temuan utama penelitian.")
    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.line_spacing = 1.0
    r_kw = p_kw.add_run("Kata Kunci: ")
    r_kw.bold = True
    p_kw.add_run("Metode, Sistem Pendukung Keputusan, Algoritma, Informatika.")

    doc.add_page_break()
    doc.add_paragraph("ABSTRACT", style='Heading 1')
    doc.add_paragraph("Abstract must be written in English with single spacing, italic style for foreign terms, containing 150-250 words describing background, objective, methodology, and key results.")
    p_kw_en = doc.add_paragraph()
    p_kw_en.paragraph_format.line_spacing = 1.0
    r_kw_en = p_kw_en.add_run("Keywords: ")
    r_kw_en.bold = True
    p_kw_en.add_run("Decision Support System, Method, Algorithm, Informatics.")

    doc.add_page_break()
    doc.add_paragraph("DAFTAR ISI", style='Heading 1')
    doc.add_paragraph("[Gunakan fitur References -> Table of Contents otomatis pada Microsoft Word]")

    doc.add_page_break()
    doc.add_paragraph("DAFTAR TABEL", style='Heading 1')
    doc.add_paragraph("[Gunakan fitur References -> Insert Table of Figures untuk Tabel]")

    doc.add_page_break()
    doc.add_paragraph("DAFTAR GAMBAR", style='Heading 1')
    doc.add_paragraph("[Gunakan fitur References -> Insert Table of Figures untuk Gambar]")

    # SECTION 3: BAB I
    sec_bab1 = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_section_geometry(sec_bab1)
    sec_bab1.header.is_linked_to_previous = False
    sec_bab1.footer.is_linked_to_previous = False
    sec_bab1.different_first_page_header_footer = True

    pg_arab = parse_xml(r'<w:pgNumType %s w:fmt="decimal" w:start="1"/>' % nsdecls('w'))
    sec_bab1._sectPr.append(pg_arab)

    p_fp_foot = sec_bab1.first_page_footer.paragraphs[0]
    p_fp_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_fp = p_fp_foot.add_run()
    r_fp.font.name = 'Times New Roman'
    r_fp.font.size = Pt(12)
    set_page_number_field(r_fp)

    p_head = sec_bab1.header.paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_head = p_head.add_run()
    r_head.font.name = 'Times New Roman'
    r_head.font.size = Pt(12)
    set_page_number_field(r_head)

    sec_bab1.first_page_header.paragraphs[0].text = ""

    doc.add_paragraph("BAB I\nPENDAHULUAN", style='Heading 1')
    doc.add_paragraph("1.1 Latar Belakang", style='Heading 2')
    p_body = doc.add_paragraph("Perkembangan teknologi informasi dan komunikasi saat ini telah memberikan dampak signifikan...")
    p_body.paragraph_format.first_line_indent = Cm(1.25)

    def add_chapter_section(bab_title, subchapters):
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        apply_section_geometry(sec)
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        sec.different_first_page_header_footer = True
        
        pg_cont = parse_xml(r'<w:pgNumType %s w:fmt="decimal"/>' % nsdecls('w'))
        sec._sectPr.append(pg_cont)

        p_fp_f = sec.first_page_footer.paragraphs[0]
        p_fp_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_f = p_fp_f.add_run()
        r_f.font.name = 'Times New Roman'
        r_f.font.size = Pt(12)
        set_page_number_field(r_f)

        p_h = sec.header.paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_h = p_h.add_run()
        r_h.font.name = 'Times New Roman'
        r_h.font.size = Pt(12)
        set_page_number_field(r_h)

        sec.first_page_header.paragraphs[0].text = ""

        doc.add_paragraph(bab_title, style='Heading 1')
        for sub_title, text_sample in subchapters:
            doc.add_paragraph(sub_title, style='Heading 2')
            p = doc.add_paragraph(text_sample)
            p.paragraph_format.first_line_indent = Cm(1.25)

    add_chapter_section("BAB II\nLANDASAN TEORI", [
        ("2.1 Tinjauan Pustaka", "Penelitian terdahulu yang relevan..."),
        ("2.2 Landasan Teori", "Teori komputasi dan metode yang digunakan...")
    ])

    add_chapter_section("BAB III\nMETODOLOGI PENELITIAN", [
        ("3.1 Waktu dan Tempat Penelitian", "Waktu dan tempat riset faktual..."),
        ("3.2 Analisis dan Desain Sistem", "Perancangan arsitektur dan basis data...")
    ])

    add_chapter_section("BAB IV\nHASIL DAN PEMBAHASAN", [
        ("4.1 Implementasi Sistem", "Antarmuka sistem nyata dan pengujian..."),
        ("4.2 Pembahasan Hasil", "Analisis hasil pengujian sistem...")
    ])

    add_chapter_section("BAB V\nKESIMPULAN DAN SARAN", [
        ("5.1 Kesimpulan", "Jawaban atas rumusan masalah..."),
        ("5.2 Saran", "Rekomendasi pengembangan sistem...")
    ])

    sec_dp = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_section_geometry(sec_dp)
    sec_dp.header.is_linked_to_previous = False
    sec_dp.footer.is_linked_to_previous = False
    sec_dp.different_first_page_header_footer = True
    
    pg_dp = parse_xml(r'<w:pgNumType %s w:fmt="decimal"/>' % nsdecls('w'))
    sec_dp._sectPr.append(pg_dp)

    p_dp_f = sec_dp.first_page_footer.paragraphs[0]
    p_dp_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_dp_f = p_dp_f.add_run()
    r_dp_f.font.name = 'Times New Roman'
    r_dp_f.font.size = Pt(12)
    set_page_number_field(r_dp_f)

    p_dp_h = sec_dp.header.paragraphs[0]
    p_dp_h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_dp_h = p_dp_h.add_run()
    r_dp_h.font.name = 'Times New Roman'
    r_dp_h.font.size = Pt(12)
    set_page_number_field(r_dp_h)

    doc.add_paragraph("DAFTAR PUSTAKA", style='Heading 1')
    p_ref_sample = doc.add_paragraph()
    p_ref_sample.paragraph_format.line_spacing = 1.0
    p_ref_sample.paragraph_format.space_after = Pt(6)
    p_ref_sample.paragraph_format.left_indent = Cm(1.25)
    p_ref_sample.paragraph_format.first_line_indent = Cm(-1.25)
    p_ref_sample.add_run("[1] Sasmita, A. U. (2021). Penerapan Metode Weighted Product pada Sistem Pendukung Keputusan Penerimaan Rekrutmen Guru. Jurnal Informatika Almuslim, 5(2), 45-52.")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Sukses membuat template skripsi: {output_path}")

if __name__ == "__main__":
    out = os.path.abspath("AI_program/working/template_skripsi_almuslim.docx")
    create_almuslim_template(out)
